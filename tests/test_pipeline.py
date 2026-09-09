"""Comprehensive test suite for the Product Label Bot pipeline."""

import unittest
import tempfile
import shutil
import time
import os
from pathlib import Path
from PIL import Image, ImageDraw

from src.config import COL_WIDTHS_DXA, LABELS_PER_PAGE, MAX_UPLOAD_IMAGE_SIDE_PX
from src.session import SessionManager, ProductItem
from src.image_ops import (
    trim_white_background,
    process_product_image,
    calculate_fit_dimensions,
)
from src.barcode_gen import is_valid_ean13, generate_barcode_image, get_barcode_encoding
from src.docx_gen import create_label_sheet
import docx


class TestImageOps(unittest.TestCase):
    def test_white_background_trim(self):
        # 300x300 white canvas with 100x100 colored square in center
        img = Image.new("RGB", (300, 300), (255, 255, 255))
        draw = ImageDraw.Draw(img)
        draw.rectangle([100, 100, 200, 200], fill=(200, 50, 50))

        cropped = trim_white_background(img, padding=5)
        # Cropped size should be approximately 101 + 2*5 = 111
        self.assertLess(cropped.size[0], 300)
        self.assertLess(cropped.size[1], 300)
        self.assertGreaterEqual(cropped.size[0], 100)
        self.assertGreaterEqual(cropped.size[1], 100)

    def test_dark_background_no_trim(self):
        # Non-white image should not be cropped
        img = Image.new("RGB", (200, 200), (30, 40, 50))
        draw = ImageDraw.Draw(img)
        draw.rectangle([50, 50, 150, 150], fill=(255, 255, 255))
        cropped = trim_white_background(img)
        self.assertEqual(cropped.size, (200, 200))

    def test_large_phone_photo_is_downscaled(self):
        temp_dir = Path(tempfile.mkdtemp())
        try:
            src = temp_dir / "large_phone_photo.jpg"
            out = temp_dir / "processed.png"
            img = Image.new("RGB", (3200, 2200), (80, 90, 100))
            img.save(src)

            width, height = process_product_image(src, out)
            self.assertTrue(out.is_file())
            self.assertLessEqual(max(width, height), MAX_UPLOAD_IMAGE_SIDE_PX)
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_docx_embedded_photo_is_compact_jpeg(self):
        temp_dir = Path(tempfile.mkdtemp())
        try:
            src = temp_dir / "large_phone_photo.jpg"
            out = temp_dir / "processed.jpg"
            img = Image.new("RGB", (3200, 2200), (80, 90, 100))
            img.save(src)

            width, height = process_product_image(
                src,
                out,
                target_max_size_px=(662, 426),
                output_format="JPEG",
                jpeg_quality=82,
            )

            self.assertTrue(out.is_file())
            self.assertLessEqual(width, 662)
            self.assertLessEqual(height, 426)
            with Image.open(out) as rendered:
                self.assertEqual(rendered.format, "JPEG")
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_phone_exif_orientation_is_normalized(self):
        temp_dir = Path(tempfile.mkdtemp())
        try:
            src = temp_dir / "rotated_phone_photo.jpg"
            out = temp_dir / "processed.png"
            img = Image.new("RGB", (60, 120), (80, 90, 100))
            exif = Image.Exif()
            exif[274] = 6  # Rotate 90 degrees
            img.save(src, exif=exif)

            width, height = process_product_image(src, out)
            self.assertTrue(out.is_file())
            self.assertGreater(width, height)
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_fit_dimensions_no_distortion(self):
        # 400x200 image (aspect 2:1) into max 50x25 mm
        w, h = calculate_fit_dimensions(400, 200, 50.0, 25.0)
        self.assertAlmostEqual(w / h, 2.0, places=1)
        self.assertLessEqual(w, 50.0)
        self.assertLessEqual(h, 25.0)

        # 100x300 image (aspect 1:3) into max 48x24 mm
        w2, h2 = calculate_fit_dimensions(100, 300, 48.0, 24.0)
        self.assertAlmostEqual(w2 / h2, 1 / 3, places=1)
        self.assertLessEqual(w2, 48.0)
        self.assertLessEqual(h2, 24.0)


class TestBarcodeGen(unittest.TestCase):
    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_code128_internal_code(self):
        out_file = self.temp_dir / "bc_code128.png"
        path, bc_type, size = generate_barcode_image("1000049", out_file)
        self.assertTrue(path.is_file())
        self.assertEqual(bc_type, "Code128")
        self.assertGreater(size[0], 100)
        self.assertGreater(size[1], 50)

    def test_ean13_detection(self):
        # Valid EAN-13
        self.assertTrue(is_valid_ean13("5901234123457"))
        # Invalid checksum
        self.assertFalse(is_valid_ean13("5901234123458"))
        # Not 13 digits
        self.assertFalse(is_valid_ean13("1000049"))

    def test_ean13_generation(self):
        out_file = self.temp_dir / "bc_ean13.png"
        path, bc_type, _ = generate_barcode_image("5901234123457", out_file)
        self.assertTrue(path.is_file())
        self.assertEqual(bc_type, "EAN-13")

    def test_ean13_invalid_checksum_fallback(self):
        out_file = self.temp_dir / "bc_fallback.png"
        # 13 digits but invalid checksum falls back to Code128 without altering digits
        path, bc_type, _ = generate_barcode_image("1234567890123", out_file)
        self.assertTrue(path.is_file())
        self.assertEqual(bc_type, "Code128")

    def test_empty_barcode_raises(self):
        with self.assertRaises(ValueError):
            generate_barcode_image("", self.temp_dir / "empty.png")

    def test_encoded_value_is_preserved_exactly(self):
        cases = [
            ("1000049", "Code128"),
            ("PRD/76:50-X SP", "Code128"),
            ("5901234123457", "EAN-13"),
            ("1234567890123", "Code128"),
        ]
        for value, expected_type in cases:
            with self.subTest(value=value):
                barcode_type, full_code = get_barcode_encoding(value)
                self.assertEqual(barcode_type, expected_type)
                self.assertEqual(full_code, value)


class TestSessionManager(unittest.TestCase):
    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp())
        self.session_base = self.temp_dir / "sessions"
        self.generated_base = self.temp_dir / "generated"
        self.mgr = SessionManager(self.session_base, self.generated_base)

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_order_preservation(self):
        session = self.mgr.get_session(12345)
        # Add products in specific order
        names = ["Prod A", "Prod B", "Prod C"]
        codes = ["001", "002", "003"]
        for n, c in zip(names, codes):
            session.add_product(n, c, "/fake/path.png")

        # Verify order
        saved_names = [p.name for p in session.products]
        self.assertEqual(saved_names, names)

    def test_remove_product(self):
        session = self.mgr.get_session(12345)
        session.add_product("First", "1", "/fake/1.png")
        session.add_product("Second", "2", "/fake/2.png")
        session.add_product("Third", "3", "/fake/3.png")

        removed = session.remove_product(2)
        self.assertIsNotNone(removed)
        self.assertEqual(removed.name, "Second")
        self.assertEqual(len(session.products), 2)
        self.assertEqual(session.products[0].name, "First")
        self.assertEqual(session.products[1].name, "Third")

    def test_update_product_fields(self):
        session = self.mgr.get_session(12345)
        old_img = session.session_dir / "old.png"
        new_img = session.session_dir / "new.png"
        old_img.write_text("old", encoding="utf-8")
        new_img.write_text("new", encoding="utf-8")
        session.add_product("Old name", "111", str(old_img))

        self.assertIsNotNone(session.update_product_name(1, "New name"))
        self.assertIsNotNone(session.update_product_barcode(1, "222"))
        self.assertIsNotNone(session.update_product_image(1, str(new_img)))

        updated = session.get_product(1)
        self.assertEqual(updated.name, "New name")
        self.assertEqual(updated.barcode, "222")
        self.assertEqual(updated.image_path, str(new_img))
        self.assertFalse(old_img.exists())
        self.assertTrue(new_img.exists())

    def test_update_product_rejects_invalid_index(self):
        session = self.mgr.get_session(12345)

        self.assertIsNone(session.get_product(1))
        self.assertIsNone(session.update_product_name(1, "Name"))
        self.assertIsNone(session.update_product_barcode(1, "123"))
        self.assertIsNone(session.update_product_image(1, "/fake/new.png"))

    def test_clear_preserves_documents(self):
        session = self.mgr.get_session(12345)
        # Create a dummy export document
        doc_file = session.session_dir / "labels_preview_20260909_120000_abc123.docx"
        pdf_file = session.session_dir / "labels_preview_20260909_120000_abc123.pdf"
        doc_file.write_text("dummy docx")
        pdf_file.write_text("dummy pdf")

        # Create a product image
        img_file = session.session_dir / "prod.png"
        img_file.write_text("dummy img")
        session.add_product("Item 1", "123", str(img_file))

        # Clear session
        session.clear()

        # Active products cleared
        self.assertEqual(len(session.products), 0)
        # Active image removed
        self.assertFalse(img_file.is_file())
        # Generated docx and pdf PRESERVED
        self.assertTrue(doc_file.is_file())
        self.assertTrue(pdf_file.is_file())

    def test_unique_export_paths_are_in_generated_dir(self):
        session = self.mgr.get_session(12345)
        first = session.new_export_docx_path("labels_preview")
        second = session.new_export_docx_path("labels_preview")

        self.assertEqual(first.parent, session.generated_dir)
        self.assertEqual(second.parent, session.generated_dir)
        self.assertNotEqual(first.name, second.name)
        self.assertRegex(first.name, r"^labels_preview_\d{8}_\d{6}_[a-f0-9]{6}\.docx$")

    def test_document_name_prefixes_export_filename(self):
        session = self.mgr.get_session(12345)
        session.set_document_name("  Santexnika Perexodlar  ")

        path = session.new_export_docx_path("product_labels")

        self.assertEqual(path.parent, session.generated_dir)
        self.assertRegex(path.name, r"^santexnika-perexodlar_product_labels_\d{8}_\d{6}_[a-f0-9]{6}\.docx$")

    def test_record_and_list_exports(self):
        session = self.mgr.get_session(12345)
        session.set_document_name("Test hujjat")
        docx_path = session.new_export_docx_path("labels_preview")
        pdf_path = docx_path.with_suffix(".pdf")
        docx_path.write_text("dummy docx", encoding="utf-8")
        pdf_path.write_text("dummy pdf", encoding="utf-8")

        record = session.record_export(
            kind="preview",
            docx_path=docx_path,
            pdf_path=pdf_path,
            product_count=3,
            page_count=1,
        )

        listed = session.list_exports()
        self.assertEqual(len(listed), 1)
        self.assertEqual(listed[0].export_id, record.export_id)
        self.assertEqual(listed[0].title, "Test hujjat")
        self.assertEqual(listed[0].kind, "preview")
        self.assertEqual(session.get_export(record.export_id).docx_path, str(docx_path.resolve()))

    def test_existing_generated_files_are_imported_to_history(self):
        session = self.mgr.get_session(12345)
        old_docx = session.generated_dir / "oldingi-hujjat_labels_preview_20260909_120000_abc123.docx"
        old_pdf = old_docx.with_suffix(".pdf")
        old_docx.write_text("dummy docx", encoding="utf-8")
        old_pdf.write_text("dummy pdf", encoding="utf-8")

        listed = session.list_exports()

        self.assertEqual(len(listed), 1)
        self.assertEqual(listed[0].title, "oldingi hujjat")
        self.assertEqual(listed[0].kind, "preview")
        self.assertEqual(listed[0].docx_path, str(old_docx.resolve()))
        self.assertEqual(listed[0].pdf_path, str(old_pdf.resolve()))
        self.assertEqual(listed[0].product_count, 0)
        self.assertEqual(listed[0].page_count, 0)

    def test_existing_legacy_session_files_are_imported_to_history(self):
        session = self.mgr.get_session(12345)
        old_docx = session.session_dir / "legacy-product_labels_20260909_120000_abc123.docx"
        old_docx.write_text("dummy docx", encoding="utf-8")

        listed = session.list_exports()

        self.assertEqual(len(listed), 1)
        self.assertEqual(listed[0].title, "legacy")
        self.assertEqual(listed[0].kind, "final")
        self.assertEqual(listed[0].docx_path, str(old_docx.resolve()))

    def test_clear_resets_active_document_but_preserves_export_history(self):
        session = self.mgr.get_session(12345)
        session.set_document_name("Mavjud hujjat")
        docx_path = session.new_export_docx_path("product_labels")
        docx_path.write_text("dummy docx", encoding="utf-8")
        record = session.record_export(
            kind="final",
            docx_path=docx_path,
            pdf_path=None,
            product_count=1,
            page_count=1,
        )

        img_file = session.session_dir / "prod.png"
        img_file.write_text("dummy img")
        session.add_product("Item 1", "123", str(img_file))
        session.clear()

        self.assertIsNone(session.document_name)
        self.assertEqual(len(session.products), 0)
        self.assertEqual(len(session.list_exports()), 1)
        self.assertEqual(session.list_exports()[0].export_id, record.export_id)

    def test_cleanup_removes_old_exports_and_stale_drafts_only(self):
        session = self.mgr.get_session(12345)
        now = time.time()
        old_ts = now - (31 * 86400)
        new_ts = now - (29 * 86400)

        old_export = session.generated_dir / "old.docx"
        new_export = session.generated_dir / "new.pdf"
        stale_draft = session.session_dir / "draft_old.jpg"
        active_draft = session.session_dir / "draft_active.jpg"

        for path in [old_export, new_export, stale_draft, active_draft]:
            path.write_text("dummy", encoding="utf-8")

        for path in [old_export, stale_draft]:
            time_tuple = (old_ts, old_ts)
            os.utime(path, time_tuple)
        for path in [new_export, active_draft]:
            time_tuple = (new_ts, new_ts)
            os.utime(path, time_tuple)

        session.current_draft["photo_path"] = str(active_draft)
        session.cleanup_expired_files(max_age_days=30, now_ts=now)

        self.assertFalse(old_export.exists())
        self.assertFalse(stale_draft.exists())
        self.assertTrue(new_export.exists())
        self.assertTrue(active_draft.exists())

    def test_cleanup_prunes_expired_export_records(self):
        session = self.mgr.get_session(12345)
        session.set_document_name("Eski hujjat")
        now = time.time()
        old_ts = now - (31 * 86400)
        old_docx = session.new_export_docx_path("product_labels")
        old_docx.write_text("dummy docx", encoding="utf-8")
        os.utime(old_docx, (old_ts, old_ts))

        record = session.record_export(
            kind="final",
            docx_path=old_docx,
            pdf_path=None,
            product_count=1,
            page_count=1,
        )
        self.assertIsNotNone(session.get_export(record.export_id))

        session.cleanup_expired_files(max_age_days=30, now_ts=now)

        self.assertFalse(old_docx.exists())
        self.assertIsNone(session.get_export(record.export_id))
        self.assertEqual(session.list_exports(), [])


class TestDocxGeneration(unittest.TestCase):
    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp())
        # Create a sample test image
        self.img_path = self.temp_dir / "test_item.png"
        img = Image.new("RGB", (200, 200), (255, 255, 255))
        d = ImageDraw.Draw(img)
        d.rectangle([40, 40, 160, 160], fill=(0, 150, 255))
        img.save(self.img_path)

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_single_product_docx(self):
        item = ProductItem("Perexod 76-50mm", "1000049", str(self.img_path))
        doc_path = self.temp_dir / "single.docx"
        created = create_label_sheet([item], doc_path)

        self.assertTrue(created.is_file())
        doc = docx.Document(created)
        self.assertEqual(len(doc.sections), 1)
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 4)
        self.assertEqual(len(table.columns), 3)

        # Check label 0 contents
        cell0 = table.cell(0, 0)
        self.assertEqual(len(cell0.paragraphs), 5)
        self.assertEqual(cell0.paragraphs[1].text, "Perexod 76-50mm")
        self.assertEqual(cell0.paragraphs[3].text, "1000049")

        # Verify fixed table layout and no 1pt line height
        from docx.oxml.ns import qn
        tblLayout = table._tbl.tblPr.find(qn("w:tblLayout"))
        self.assertIsNotNone(tblLayout)
        self.assertEqual(tblLayout.get(qn("w:type")), "fixed")
        tblW = table._tbl.tblPr.find(qn("w:tblW"))
        self.assertIsNotNone(tblW)
        self.assertEqual(tblW.get(qn("w:type")), "dxa")
        self.assertEqual(int(tblW.get(qn("w:w"))), sum(COL_WIDTHS_DXA))
        grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
        self.assertEqual([int(col.get(qn("w:w"))) for col in grid_cols], COL_WIDTHS_DXA)

        # Verify no paragraph has fatal 1pt line spacing
        for p in cell0.paragraphs:
            if p._p.pPr is not None:
                spacing = p._p.pPr.find(qn("w:spacing"))
                if spacing is not None:
                    self.assertFalse(
                        spacing.get(qn("w:lineRule")) == "exact" and spacing.get(qn("w:line")) == "20",
                        "Fatal 1pt line spacing detected!"
                    )

    def test_special_characters_barcode(self):
        # Barcodes with slashes, colons, spaces, dashes (Code128 supported)
        item = ProductItem("Special Item", "PRD/76:50-X SP", str(self.img_path))
        doc_path = self.temp_dir / "special_bc.docx"
        created = create_label_sheet([item], doc_path)
        self.assertTrue(created.is_file())
        doc = docx.Document(created)
        cell0 = doc.tables[0].cell(0, 0)
        self.assertEqual(cell0.paragraphs[3].text, "PRD/76:50-X SP")

    def test_multi_page_pagination(self):
        # 13 items should produce 2 pages (12 on page 1, 1 on page 2)
        items = [
            ProductItem(f"Product {i}", f"1000{i:03d}", str(self.img_path))
            for i in range(13)
        ]
        doc_path = self.temp_dir / "thirteen.docx"
        created = create_label_sheet(items, doc_path)

        doc = docx.Document(created)
        self.assertEqual(len(doc.sections), 2)
        self.assertEqual(len(doc.tables), 2)

        # Page 1 has 12 items
        self.assertEqual(doc.tables[0].cell(3, 2).paragraphs[1].text, "Product 11")
        # Page 2 has 13th item (index 12) at cell(0, 0)
        self.assertEqual(doc.tables[1].cell(0, 0).paragraphs[1].text, "Product 12")
        # Remaining cells on page 2 are empty for cutting
        self.assertEqual(doc.tables[1].cell(0, 1).paragraphs[0].text, "")


class TestBotEscaping(unittest.TestCase):
    def test_markdown_escaping(self):
        from src.bot import esc
        raw = "Item_Name*with[brackets]and`code`_100%"
        escaped = esc(raw)
        self.assertIn(r"\*with", escaped)
        self.assertIn(r"\[brackets", escaped)
        self.assertIn(r"\_", escaped)

    def test_safe_image_extension(self):
        from src.bot import safe_image_extension
        self.assertEqual(safe_image_extension("photo.PNG", "image/png"), ".png")
        self.assertEqual(safe_image_extension("photo.jpeg", None), ".jpeg")
        self.assertEqual(safe_image_extension("upload.txt", "image/jpeg"), ".jpg")
        self.assertEqual(safe_image_extension("bad.png/evil", "image/png"), ".png")
        self.assertEqual(safe_image_extension("vector.svg", None), ".png")

    def test_webhook_url_helpers(self):
        from src.bot import build_webhook_public_url, normalize_webhook_path
        self.assertEqual(normalize_webhook_path("tg/secret"), "/tg/secret")
        self.assertEqual(normalize_webhook_path("/tg/secret"), "/tg/secret")
        self.assertEqual(
            build_webhook_public_url("https://labels.example.com/", "tg/secret"),
            "https://labels.example.com/tg/secret",
        )

    def test_product_list_has_edit_buttons(self):
        from src.bot import build_product_edit_keyboard, build_products_keyboard, format_products_list

        temp_dir = Path(tempfile.mkdtemp())
        try:
            mgr = SessionManager(temp_dir / "sessions", temp_dir / "generated")
            session = mgr.get_session(12345)
            img_file = session.session_dir / "item.png"
            img_file.write_text("img", encoding="utf-8")
            session.set_document_name("Test")
            session.add_product("Scanner", "1000023", str(img_file))
            session.add_product("Perexod 76-50mm", "1000049", str(img_file))

            list_text = format_products_list(session)
            self.assertIn("/edit 2", list_text)
            self.assertIn("Perexod", list_text)

            keyboard = build_products_keyboard(session)
            callback_data = [
                button.callback_data
                for row in keyboard.inline_keyboard
                for button in row
            ]
            self.assertIn("edit:1", callback_data)
            self.assertIn("edit:2", callback_data)
            self.assertIn("btn_add", callback_data)

            edit_keyboard = build_product_edit_keyboard(2)
            edit_callbacks = [
                button.callback_data
                for row in edit_keyboard.inline_keyboard
                for button in row
            ]
            self.assertEqual(
                edit_callbacks,
                ["edit_photo:2", "edit_name:2", "edit_barcode:2", "remove:2", "btn_list"],
            )
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_reply_keyboard_is_optional_not_persistent(self):
        from src.bot import MAIN_REPLY_KEYBOARD

        self.assertTrue(MAIN_REPLY_KEYBOARD.one_time_keyboard)
        self.assertFalse(MAIN_REPLY_KEYBOARD.is_persistent)

    def test_large_document_chunks_preserve_order(self):
        from src.bot import chunk_products_for_document_parts, next_smaller_docx_part_size, safe_docx_part_size

        products = list(range(25))
        chunks = chunk_products_for_document_parts(products, 12)

        self.assertEqual(chunks, [list(range(12)), list(range(12, 24)), [24]])
        self.assertEqual(safe_docx_part_size(100), 72)
        self.assertEqual(next_smaller_docx_part_size(72), 36)
        self.assertEqual(next_smaller_docx_part_size(12), 6)


if __name__ == "__main__":
    unittest.main()
