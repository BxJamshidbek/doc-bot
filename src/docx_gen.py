"""DOCX A4 12-label (3x4 grid) sheet generator."""

import hashlib
import math
import shutil
import tempfile
import uuid
from pathlib import Path
from typing import List, Sequence

import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_SECTION_START
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image, ImageDraw

from src.config import (
    PAGE_WIDTH_MM,
    PAGE_HEIGHT_MM,
    PAGE_MARGIN_TOP_MM,
    PAGE_MARGIN_BOTTOM_MM,
    PAGE_MARGIN_LEFT_MM,
    PAGE_MARGIN_RIGHT_MM,
    GRID_COLS,
    GRID_ROWS,
    LABELS_PER_PAGE,
    COL_WIDTHS_MM,
    ROW_HEIGHT_MM,
    ROW_HEIGHT_DXA,
    MAX_PHOTO_WIDTH_MM,
    MAX_PHOTO_HEIGHT_MM,
    MAX_BARCODE_WIDTH_MM,
    MAX_BARCODE_HEIGHT_MM,
    FONT_NAME,
    FONT_SIZE_NAME_PT,
    FONT_SIZE_CODE_PT,
    BORDER_COLOR_OUTER,
    BORDER_COLOR_INNER,
)
from src.image_ops import process_product_image, calculate_fit_dimensions
from src.barcode_gen import generate_barcode_image
from src.session import ProductItem

EMPTY_CELL_SPACER_WIDTH_MM = 62.0
EMPTY_CELL_SPACER_HEIGHT_MM = 0.2


def _remove_child_elements(parent, tag_name: str) -> None:
    """Removes all direct OOXML children with the given qualified tag name."""
    for child in list(parent.findall(qn(tag_name))):
        parent.remove(child)


def _apply_section_page_setup(section) -> None:
    """Applies A4 page dimensions and margins to a Word section."""
    section.page_width = Mm(PAGE_WIDTH_MM)
    section.page_height = Mm(PAGE_HEIGHT_MM)
    section.top_margin = Mm(PAGE_MARGIN_TOP_MM)
    section.bottom_margin = Mm(PAGE_MARGIN_BOTTOM_MM)
    section.left_margin = Mm(PAGE_MARGIN_LEFT_MM)
    section.right_margin = Mm(PAGE_MARGIN_RIGHT_MM)


def _setup_table_properties(table) -> None:
    """Applies fixed layout and light gray cutting guide borders to the table."""
    table.autofit = False
    tblPr = table._tbl.tblPr

    total_width_dxa = sum(int(round(w * 56.6929)) for w in COL_WIDTHS_MM)

    # Replace python-docx's default auto width in-place. Some DOCX preview
    # renderers are sensitive to table-property ordering, so keep w:tblW at
    # its original position when python-docx already created it.
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblPr.insert(0, parse_xml(f'<w:tblW {nsdecls("w")} w:w="{total_width_dxa}" w:type="dxa"/>'))
    else:
        tblW.set(qn("w:w"), str(total_width_dxa))
        tblW.set(qn("w:type"), "dxa")

    # Fixed table layout in OOXML
    _remove_child_elements(tblPr, "w:tblLayout")
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))

    # Light gray borders for cutting guides
    _remove_child_elements(tblPr, "w:tblBorders")
    borders_xml = f"""
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR_OUTER}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR_OUTER}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR_OUTER}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR_OUTER}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR_INNER}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR_INNER}"/>
        </w:tblBorders>
    """
    tblPr.append(parse_xml(borders_xml))

    # Explicit column widths in tblGrid. Always replace the default grid so
    # renderers cannot keep stale auto-fit measurements from table creation.
    for tblGrid in list(table._tbl.findall(qn("w:tblGrid"))):
        table._tbl.remove(tblGrid)

    grid_cols_xml = "".join(
        f'<w:gridCol {nsdecls("w")} w:w="{int(round(w * 56.6929))}"/>'
        for w in COL_WIDTHS_MM
    )
    grid_element = parse_xml(f'<w:tblGrid {nsdecls("w")}>{grid_cols_xml}</w:tblGrid>')
    table._tbl.insert(1, grid_element)

    for idx, width_mm in enumerate(COL_WIDTHS_MM):
        table.columns[idx].width = Mm(width_mm)


def _set_cell_properties(cell, width_mm: float) -> None:
    """Configures cell width, margins, and vertical alignment with explicit OOXML."""
    cell.width = Mm(width_mm)
    tcPr = cell._tc.get_or_add_tcPr()

    # Explicit cell width in OOXML w:tcW
    width_dxa = int(round(width_mm * 56.6929))
    _remove_child_elements(tcPr, "w:tcW")
    tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>'))

    # Center vertical alignment
    _remove_child_elements(tcPr, "w:vAlign")
    tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

    # Cell internal padding (dxa)
    _remove_child_elements(tcPr, "w:tcMar")
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="40" w:type="dxa"/>'
        f'  <w:bottom w:w="40" w:type="dxa"/>'
        f'  <w:left w:w="80" w:type="dxa"/>'
        f'  <w:right w:w="80" w:type="dxa"/>'
        f'</w:tcMar>'
    ))


def _empty_cell_spacer_path(temp_dir: Path) -> Path:
    """Returns a reusable transparent spacer image for empty label cells."""
    spacer_path = temp_dir / "empty_cell_width_spacer.png"
    if not spacer_path.is_file():
        Image.new("RGBA", (20, 1), (255, 255, 255, 0)).save(spacer_path)
    return spacer_path


def _add_width_anchor_to_paragraph(paragraph, temp_dir: Path, width_mm: float) -> None:
    """Adds an invisible inline image that forces preview renderers to keep cell width."""
    spacer_width = min(EMPTY_CELL_SPACER_WIDTH_MM, max(1.0, width_mm - 1.0))
    paragraph.add_run().add_picture(
        str(_empty_cell_spacer_path(temp_dir)),
        width=Mm(spacer_width),
        height=Mm(EMPTY_CELL_SPACER_HEIGHT_MM),
    )


def _append_width_anchor_paragraph(cell, temp_dir: Path, width_mm: float) -> None:
    """Adds a zero-visual-impact width anchor at the end of a populated cell."""
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    _add_width_anchor_to_paragraph(p, temp_dir, width_mm)


def _populate_empty_label_cell(cell, temp_dir: Path, width_mm: float) -> None:
    """
    Keeps empty cells visually blank while forcing preview renderers to retain
    the intended column width. Some DOCX previews collapse completely empty
    columns even with fixed table widths.
    """
    p = cell.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0

    _add_width_anchor_to_paragraph(p, temp_dir, width_mm)



def _populate_label_cell(
    cell,
    product: ProductItem,
    temp_dir: Path,
    width_mm: float,
) -> None:
    """
    Fills a single cell with product photo, name, barcode, and barcode number.
    Uses normal proportional line spacing (line_spacing = 1.0) so images and
    text render with full visibility and zero clipping.
    """
    # 1. Product Photo (Inline, centered, normal line spacing)
    p_img = cell.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(0)
    p_img.paragraph_format.space_after = Pt(1)
    p_img.paragraph_format.line_spacing = 1.0

    photo_path = Path(product.image_path)
    # Safe temp filename using uuid4 to support any original filename
    processed_photo = temp_dir / f"proc_{uuid.uuid4().hex}.png"

    try:
        if photo_path.is_file():
            img_w, img_h = process_product_image(photo_path, processed_photo)
        else:
            placeholder = Image.new("RGB", (200, 200), (245, 245, 245))
            d = ImageDraw.Draw(placeholder)
            d.text((50, 90), "No Image", fill=(120, 120, 120))
            placeholder.save(processed_photo)
            img_w, img_h = 200, 200
    except Exception:
        placeholder = Image.new("RGB", (200, 200), (245, 245, 245))
        d = ImageDraw.Draw(placeholder)
        d.text((50, 90), "No Image", fill=(120, 120, 120))
        placeholder.save(processed_photo)
        img_w, img_h = 200, 200

    fit_w, fit_h = calculate_fit_dimensions(
        img_w, img_h, MAX_PHOTO_WIDTH_MM, MAX_PHOTO_HEIGHT_MM
    )
    r_img = p_img.add_run()
    r_img.add_picture(str(processed_photo), width=Mm(fit_w), height=Mm(fit_h))

    # 2. Product Name (Bold, Centered, normal line spacing)
    p_name = cell.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    p_name.paragraph_format.line_spacing = 1.0

    r_name = p_name.add_run(product.name)
    r_name.font.name = FONT_NAME
    name_len = len(product.name)
    if name_len <= 24:
        name_size = FONT_SIZE_NAME_PT
    elif name_len <= 38:
        name_size = FONT_SIZE_NAME_PT - 0.6
    elif name_len <= 54:
        name_size = FONT_SIZE_NAME_PT - 1.2
    else:
        name_size = FONT_SIZE_NAME_PT - 2.0
    r_name.font.size = Pt(name_size)
    r_name.bold = True
    r_name.font.color.rgb = RGBColor(20, 20, 20)

    # 3. Barcode Image (Inline, centered, normal line spacing)
    p_bc = cell.add_paragraph()
    p_bc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bc.paragraph_format.space_before = Pt(0)
    p_bc.paragraph_format.space_after = Pt(0)
    p_bc.paragraph_format.line_spacing = 1.0

    # Safe temp barcode filename using sha256 to support slashes, colons, spaces, etc.
    bc_hash = hashlib.sha256(product.barcode.encode("utf-8", errors="replace")).hexdigest()[:16]
    bc_output_path = temp_dir / f"bc_{bc_hash}_{uuid.uuid4().hex[:6]}.png"
    bc_path, _, bc_size = generate_barcode_image(product.barcode, bc_output_path)

    bc_w, bc_h = calculate_fit_dimensions(
        bc_size[0], bc_size[1], MAX_BARCODE_WIDTH_MM, MAX_BARCODE_HEIGHT_MM
    )
    r_bc = p_bc.add_run()
    r_bc.add_picture(str(bc_path), width=Mm(bc_w), height=Mm(bc_h))

    # 4. Barcode Code Text (Centered, Clear, normal line spacing)
    p_code = cell.add_paragraph()
    p_code.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_code.paragraph_format.space_before = Pt(0)
    p_code.paragraph_format.space_after = Pt(0)
    p_code.paragraph_format.line_spacing = 1.0

    r_code = p_code.add_run(product.barcode)
    r_code.font.name = FONT_NAME
    r_code.font.size = Pt(FONT_SIZE_CODE_PT)
    r_code.bold = True
    r_code.font.color.rgb = RGBColor(30, 30, 30)

    _append_width_anchor_paragraph(cell, temp_dir, width_mm)


def create_label_sheet(
    products: Sequence[ProductItem],
    output_path: str | Path
) -> Path:
    """
    Creates an A4 document containing 12 labels per page (3 cols x 4 rows).
    Preserves the exact order of products.
    Returns the resolved output Path.
    """
    if not products:
        raise ValueError("Cannot create label sheet with empty product list.")

    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    temp_render_dir = Path(tempfile.mkdtemp(prefix="label_render_"))

    try:
        doc = docx.Document()
        total_products = len(products)
        num_pages = math.ceil(total_products / LABELS_PER_PAGE)

        for page_idx in range(num_pages):
            if page_idx == 0:
                sec = doc.sections[0]
            else:
                sec = doc.add_section(WD_SECTION_START.NEW_PAGE)

            _apply_section_page_setup(sec)

            table = doc.add_table(rows=GRID_ROWS, cols=GRID_COLS)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            _setup_table_properties(table)

            # Configure rows with exact OOXML trHeight and cantSplit
            for row in table.rows:
                trPr = row._tr.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                trPr.append(parse_xml(f'<w:trHeight {nsdecls("w")} w:val="{ROW_HEIGHT_DXA}" w:hRule="exact"/>'))
                row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
                row.height = Mm(ROW_HEIGHT_MM)

            # Fill 12 cells for current page
            page_start = page_idx * LABELS_PER_PAGE
            for i in range(LABELS_PER_PAGE):
                row_idx = i // GRID_COLS
                col_idx = i % GRID_COLS
                cell = table.cell(row_idx, col_idx)
                _set_cell_properties(cell, COL_WIDTHS_MM[col_idx])

                prod_idx = page_start + i
                if prod_idx < total_products:
                    _populate_label_cell(cell, products[prod_idx], temp_render_dir, COL_WIDTHS_MM[col_idx])
                else:
                    # Empty cell retains cutting borders
                    _populate_empty_label_cell(cell, temp_render_dir, COL_WIDTHS_MM[col_idx])

        doc.save(str(output_path))
        return output_path
    finally:
        shutil.rmtree(temp_render_dir, ignore_errors=True)
