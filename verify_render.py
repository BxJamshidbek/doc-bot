"""Verification script for DOCX & PDF rendering: 1, 12, and 13 products with PDF page count verification."""

import shutil
import subprocess
from pathlib import Path
from PIL import Image
import docx
from docx.oxml.ns import qn
from pypdf import PdfReader

from src.session import ProductItem
from src.docx_gen import create_label_sheet
from src.pdf_converter import convert_docx_to_pdf
from test_sample import generate_sample_photo


def verify_docx_ooxml_structure(docx_path: Path, expected_pages: int):
    """Inspects the Word XML structure to ensure fixed layout and absence of 1pt line spacing."""
    doc = docx.Document(docx_path)
    assert len(doc.sections) == expected_pages, f"Expected {expected_pages} sections, found {len(doc.sections)}"
    assert len(doc.tables) == expected_pages, f"Expected {expected_pages} tables, found {len(doc.tables)}"

    total_cells_checked = 0
    for t_idx, table in enumerate(doc.tables):
        # 1. Verify w:tblLayout w:type="fixed"
        tblPr = table._tbl.tblPr
        tblLayout = tblPr.find(qn("w:tblLayout"))
        assert tblLayout is not None, f"Table {t_idx} is missing w:tblLayout"
        assert tblLayout.get(qn("w:type")) == "fixed", f"Table {t_idx} w:tblLayout is not 'fixed'"

        # 2. Verify rows have exact height and cantSplit
        assert len(table.rows) == 4, f"Table {t_idx} rows count != 4"
        for r_idx, row in enumerate(table.rows):
            trPr = row._tr.trPr
            assert trPr.find(qn("w:cantSplit")) is not None, f"Table {t_idx} Row {r_idx} missing w:cantSplit"
            trHeight = trPr.find(qn("w:trHeight"))
            assert trHeight is not None, f"Table {t_idx} Row {r_idx} missing w:trHeight"
            assert trHeight.get(qn("w:hRule")) == "exact", f"Table {t_idx} Row {r_idx} hRule != exact"

            # 3. Verify cells have w:tcW and paragraphs do not use 1pt line height
            assert len(row.cells) == 3, f"Table {t_idx} Row {r_idx} columns != 3"
            for c_idx, cell in enumerate(row.cells):
                tcPr = cell._tc.tcPr
                assert tcPr.find(qn("w:tcW")) is not None, f"Cell missing w:tcW"

                for p_idx, p in enumerate(cell.paragraphs):
                    pPr = p._p.pPr
                    if pPr is not None:
                        spacing = pPr.find(qn("w:spacing"))
                        if spacing is not None:
                            line_rule = spacing.get(qn("w:lineRule"))
                            line_val = spacing.get(qn("w:line"))
                            # Ensure no exact 1pt (20 dxa) line spacing
                            assert not (line_rule == "exact" and line_val == "20"), (
                                f"Found fatal 1pt line spacing in Table {t_idx}, Row {r_idx}, Cell {c_idx}, P {p_idx}!"
                            )
                total_cells_checked += 1

    print(f"  ✓ OOXML check passed: fixed layout, exact heights, cantSplit, and no 1pt line clipping ({total_cells_checked} cells verified).")


def verify_pdf_page_count(pdf_path: Path, expected_pages: int):
    """Verifies the actual page count of the generated PDF file using pypdf."""
    assert pdf_path.is_file(), f"PDF file not found: {pdf_path}"
    assert pdf_path.stat().st_size > 0, f"PDF file is empty: {pdf_path}"

    reader = PdfReader(str(pdf_path))
    actual_pages = len(reader.pages)
    assert actual_pages == expected_pages, (
        f"PDF Page count mismatch for {pdf_path.name}: expected {expected_pages}, found {actual_pages} page(s)!"
    )
    print(f"  ✓ PDF Page count verified: {pdf_path.name} has exactly {actual_pages} page(s) (matches expected {expected_pages}).")


def _assert_thumbnail_has_content(thumb_path: Path, label: str) -> Path:
    """Checks that a rendered page thumbnail is not blank."""
    im = Image.open(thumb_path)
    extrema = im.convert("L").getextrema()
    assert extrema[0] < 50, f"Thumbnail appears empty or all-white! Min pixel value: {extrema[0]}"
    layout_bbox = im.convert("L").point(lambda px: 0 if px >= 248 else 255).getbbox()
    assert layout_bbox is not None, "Rendered page has no detectable label/table content."
    content_width = layout_bbox[2] - layout_bbox[0]
    width_ratio = content_width / im.size[0]
    assert width_ratio >= 0.80, (
        f"Rendered label table is too narrow: content spans {width_ratio:.1%} of page width "
        f"(bbox={layout_bbox}, image_width={im.size[0]})."
    )
    print(f"  ✓ {label}: {thumb_path.name} ({im.size[0]}x{im.size[1]} px, dark pixels confirmed).")
    return thumb_path


def run_visual_thumbnail_check(docx_path: Path, out_dir: Path, pdf_path: Path):
    """Renders page 1 to PNG and checks that label layout spans the page."""
    checked_thumbnail = None
    pdftoppm = shutil.which("pdftoppm")
    if pdftoppm:
        out_prefix = out_dir / f"{pdf_path.stem}_page1"
        thumb_path = out_prefix.with_suffix(".png")
        thumb_path.unlink(missing_ok=True)
        subprocess.run(
            [pdftoppm, "-png", "-singlefile", "-f", "1", "-r", "144", str(pdf_path), str(out_prefix)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            check=True,
        )
        if thumb_path.is_file():
            checked_thumbnail = _assert_thumbnail_has_content(thumb_path, "Poppler PDF visual render")

    qlmanage = shutil.which("qlmanage")
    if qlmanage:
        thumb_path = out_dir / f"{docx_path.name}.png"
        thumb_path.unlink(missing_ok=True)
        subprocess.run(
            [qlmanage, "-t", "-s", "1200", "-o", str(out_dir), str(docx_path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            check=False,
        )
        if thumb_path.is_file():
            checked_thumbnail = _assert_thumbnail_has_content(thumb_path, "QuickLook DOCX visual thumbnail")

    if checked_thumbnail is not None:
        return checked_thumbnail

    raise RuntimeError("No visual renderer found. Install poppler-utils on Linux or use macOS QuickLook.")


def main():
    test_dir = Path("render_verification_outputs").resolve()
    test_dir.mkdir(parents=True, exist_ok=True)
    sample_img = test_dir / "sample_test_prod.png"
    generate_sample_photo(sample_img)

    # Required test cases: 1, 12, and 13 products
    test_cases = [
        (1, "render_1_product.docx", "render_1_product.pdf", 1),
        (12, "render_12_products.docx", "render_12_products.pdf", 1),
        (13, "render_13_products.docx", "render_13_products.pdf", 2),
    ]

    sample_barcodes = [
        "1000049",
        "5901234123457",
        "PRD/76:50-X",
        "CODE 128 SP",
        "ITEM#004",
        "1000055",
        "1000056",
        "1000057",
        "1000058",
        "1000059",
        "1000060",
        "1000061",
        "1000062",
    ]

    for count, docx_name, pdf_name, expected_pages in test_cases:
        print(f"\n--- Testing {count} Product(s) -> Expected {expected_pages} Page(s) ---")
        docx_path = test_dir / docx_name

        products = [
            ProductItem(
                name=f"Product {i+1} (76-50mm)",
                barcode=sample_barcodes[i % len(sample_barcodes)],
                image_path=str(sample_img),
            )
            for i in range(count)
        ]

        # 1. Generate DOCX
        create_label_sheet(products, docx_path)
        print(f"  ✓ DOCX Generated: {docx_path.name} ({docx_path.stat().st_size} bytes)")

        # 2. Verify OOXML structure
        verify_docx_ooxml_structure(docx_path, expected_pages)

        # 3. Convert DOCX to PDF
        pdf_path, pdf_note = convert_docx_to_pdf(docx_path, test_dir)
        if pdf_path and pdf_path.is_file():
            print(f"  ✓ PDF Generated: {pdf_path.name} ({pdf_path.stat().st_size} bytes)")
            # 4. Verify actual PDF page count
            verify_pdf_page_count(pdf_path, expected_pages)
        else:
            raise RuntimeError(f"Failed to generate PDF for {docx_name}: {pdf_note}")

        # 5. Visual page-render check
        run_visual_thumbnail_check(docx_path, test_dir, pdf_path)

    print("\n" + "=" * 60)
    print("✅ All DOCX and PDF render verifications passed successfully!")
    print("   • render_1_product.pdf  : exactly 1 page")
    print("   • render_12_products.pdf : exactly 1 page")
    print("   • render_13_products.pdf : exactly 2 pages")
    print("=" * 60)


if __name__ == "__main__":
    main()
